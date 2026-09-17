# -*- coding: utf-8 -*-
"""Convierte un documento Markdown de SIGAD a Word (.docx) con portada.

Uso:  python tools/generate_docx.py <origen.md> <destino.docx> [título de portada]
Formato: párrafos justificados con viñetas para los elementos de lista;
imágenes en notación Markdown (![pie](ruta)) incrustadas con su pie; portada
institucional.
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ACCENT = RGBColor(0x1F, 0x3B, 0x73)
GRAY = RGBColor(0x55, 0x55, 0x55)

BULLET_RE = re.compile(r"^(\s*)[-*]\s+")
LABEL_RE = re.compile(r"^[A-Z]{1,4}\d{1,3}\.\s")
IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")

COVER = {
    "sigad": "SIGAD",
    "tagline": (
        "Sistema web inteligente para la gestión, clasificación, análisis "
        "y consulta de documentos empresariales."
    ),
    "curso": "Desarrollo de Aplicaciones Empresariales",
    "titulo": "Documento de Diseño",
    "universidad": "Universitaria Tecnológica de Santander",
    "docente": "Docente: Wilson Castaño Galviz",
    "integrante": "Integrante: Javier Andrés Núñez Sánchez",
}


def cover_paragraph(doc, text, size, bold=False, italic=False, color=None, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p


def build_cover(doc, titulo):
    for _ in range(3):
        doc.add_paragraph()
    cover_paragraph(doc, COVER["sigad"], 40, bold=True, color=ACCENT, space_after=10)
    cover_paragraph(doc, COVER["tagline"], 13, italic=True, space_after=26)
    cover_paragraph(doc, COVER["curso"], 16, space_after=4)
    cover_paragraph(doc, titulo or COVER["titulo"], 18, bold=True, color=ACCENT, space_after=8)
    for _ in range(5):
        doc.add_paragraph()
    cover_paragraph(doc, COVER["universidad"], 14, space_after=4)
    cover_paragraph(doc, COVER["docente"], 12, space_after=2)
    cover_paragraph(doc, COVER["integrante"], 12, space_after=2)
    doc.add_page_break()


def add_image(doc, path, alt, base_dir):
    full = path if os.path.isabs(path) else os.path.join(base_dir, path)
    if not os.path.exists(full):
        body_paragraph(doc, "[Imagen no encontrada: %s]" % path)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    from PIL import Image as PILImage
    wpx, hpx = PILImage.open(full).size
    max_w_cm, max_h_cm = 16.2, 17.0
    page_w_cm, page_h_cm = 21.0, 29.7
    avail_w = page_w_cm - 5.6
    avail_h = page_h_cm - 5.0
    w_cm = min(max_w_cm, avail_w)
    h_cm = w_cm * hpx / wpx
    if h_cm > max_h_cm or h_cm > avail_h:
        h_cm = min(max_h_cm, avail_h)
        w_cm = h_cm * wpx / hpx
    run.add_picture(full, width=Cm(w_cm), height=Cm(h_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cr = cap.add_run(alt)
    cr.italic = True
    cr.font.size = Pt(10)
    cr.font.color.rgb = GRAY


def body_paragraph(doc, text, heading=False, bullet=False, indent=0):
    if bullet:
        style = "List Bullet" if indent == 0 else "List Bullet 2"
        p = doc.add_paragraph(style=style)
    else:
        p = doc.add_paragraph(style="Heading 1" if heading else "Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    if bullet:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if heading:
        p.paragraph_format.space_before = Pt(18)
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.color.rgb = ACCENT
        run.bold = True
    else:
        run = p.add_run(text)
        run.font.size = Pt(11.5)
    return p


def main(src, out, titulo=None):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.15

    build_cover(doc, titulo)
    base_dir = os.path.dirname(os.path.abspath(src))

    with open(src, encoding="utf-8") as fh:
        for raw in fh:
            line = raw.rstrip()
            stripped = line.strip()
            if not stripped:
                continue
            im = IMG_RE.match(stripped)
            if im:
                add_image(doc, im.group(2), im.group(1), base_dir)
                continue
            if stripped.startswith("# "):
                p = doc.add_paragraph(style="Title")
                run = p.add_run(stripped[2:])
                run.font.size = Pt(20)
                run.font.color.rgb = ACCENT
                run.bold = True
                p.paragraph_format.space_after = Pt(4)
            elif stripped.startswith("## "):
                body_paragraph(doc, stripped[3:], heading=True)
            else:
                m = BULLET_RE.match(line or "")
                if m:
                    indent = 0 if m.group(1) == "" else 1
                    body_paragraph(doc, re.sub(BULLET_RE, "", line).strip(), bullet=True, indent=indent)
                elif LABEL_RE.match(stripped):
                    body_paragraph(doc, stripped, bullet=True)
                else:
                    body_paragraph(doc, stripped)

    doc.save(out)
    print("OK ->", out)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python tools/generate_docx.py <origen.md> <destino.docx> [título de portada]")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)