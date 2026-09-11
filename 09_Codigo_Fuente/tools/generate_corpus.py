# -*- coding: utf-8 -*-
"""Generador del repositorio de pruebas SIGAD.

Crea 33 documentos sintéticos (sin datos personales reales) distribuidos en
4 categorías y 3 formatos (TXT, PDF, DOCX), que son los que el sistema
procesa, clasifica, resume y consulta.

Uso:
    python tools/generate_corpus.py [directorio_destino]
"""
import os
import sys
import csv
import datetime

os.chdir(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

EMISORES = {
    "emisores_fact": [
        ("Servicios Contables Andes S.A.S.", "901.234.567-8", "facturacion@contablesandes.com", "(607) 655-1234"),
        ("TecnoSuministros del Oriente S.A.S.", "901.222.334-1", "ventas@tecnosum.com.co", "(607) 634-7788"),
        ("Mantenimientos Industriales JR Ltda.", "900.555.666-7", "servicio@jrindustrial.com", "+57 607 689 4545"),
        ("Distribuidora El Puerto S.A.S.", "900.318.254-3", "cartera@elpuerto.co", "(607) 697-2415"),
        ("Logistica Ciclo Verde S.A.S.", "901.098.112-5", "comercial@ciclove.com", "607 210 3344"),
    ],
    "clientes_fact": [
        ("Comercializadora El Roble S.A.S.", "900.876.543-2"),
        ("Colegio San Rafael", "800.112.233-4"),
        ("Panaderia La Espiga", "890.900.111-5"),
        ("Clinica Vida Plena", "900.661.123-9"),
        ("Constructora Horizonte Ltda.", "804.000.987-6"),
        ("Textiles del Oriente S.A.S.", "901.444.567-8"),
    ],
    "contrato": [
        ("Tecnologia Andina S.A.S.", "900.321.654-9", "(607) 689-4545", "legal@tecnologiaandina.com"),
        ("Inversiones Roble Verde S.A.S.", "901.456.789-0", "(607) 642-1900", "juridica@robleverde.com"),
        ("Corporacion Biosalud SG", "890.223.445-6", "(607) 645-8820", "contratos@biosalud.net"),
        ("Energia Industrial del Oriente S.A.", "890.111.332-4", "607 634 9988", "adquisiciones@eio.co"),
    ],
    "personas": ["Ing. Laura Martinez Rueda", "Sr. Ricardo Gomez Pena", "Sra. Marcela Torres Quinones",
                 "Dr. Andres Camargo Silva", "Abg. Felipe Ruiz Guzman", "Sra. Paula Andrea Nieto",
                 "Sr. Camilo Andres Rueda", "Ing. Oscar Villamizar", "Sra. Diana Salazar Ortiz",
                 "Dr. Samuel Restrepo"],
    "orgs_comun": ["Distribuidora El Puerto S.A.S.", "Logistica Central S.A.S.", "Comercializadora El Roble S.A.S."],
}

MONTHS_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio",
             "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def fecha_es(dt):
    return "%d de %s de %d" % (dt.day, MONTHS_ES[dt.month - 1], dt.year)


def build_factura(i):
    emisor, nit_e, mail, tel = EMISORES["emisores_fact"][i % 5]
    cliente, nit_c = EMISORES["clientes_fact"][i % 6]
    exp = datetime.date(2024, (i % 3) + 1, 5 + i)
    ven = exp + datetime.timedelta(days=45)
    units = 1 + (i % 4)
    unit_val = [450000, 1840000, 150000, 820000, 120000][i % 5] + (i * 17000)
    sub, iva, total = units * unit_val, 0, 0
    perc = 0.19
    iva = round(sub * perc)
    total = sub + iva
    concepto = ["honorarios contables y tributarios del periodo",
                "suministro de equipos y repuestos tecnologicos",
                "mantenimiento preventivo y predictivo de maquinaria",
                "servicios de aseo y cafeteria corporativa",
                "capacitacion en normatividad laboral al personal"][i % 5]
    return ("FACTURA ELECTRONICA DE VENTA No. FE-%03d\n"
            "%s - NIT %s\n"
            "Telefono: %s - Correo: %s\n"
            "Cliente: %s - NIT %s\n"
            "Fecha de expedicion: %s\n"
            "Fecha de vencimiento: %s\n"
            "Concepto: %s - Cantidad: %d unidades\n"
            "Valor unitario: $%s - Subtotal: $%s\n"
            "IVA (19%%): $%s\n"
            "TOTAL A PAGAR: COP $%s\n"
            "Forma de pago: %s\n"
            "Resolucion de facturacion DIAN No. 18764000%05d\n\n"
            "La presente factura constituye titulo valor conforme al articulo 774 del Codigo de Comercio.\n"
            ) % (
        i + 1, emisor, nit_e, tel, mail, cliente, nit_c,
        fecha_es(exp), fecha_es(ven), concepto, units,
        "{:,}".format(unit_val), "{:,}".format(sub), "{:,}".format(iva), "{:,}".format(total),
        ["transferencia electronica", "consignacion bancaria", "efectivo", "cheque", "pago a credito 30 dias"][i % 5],
        i)


def build_contrato(i):
    tipo = ["arrendamiento de inmueble para uso de oficina",
            "prestacion de servicios profesionales",
            "confidencialidad sobre informacion tecnica",
            "mantenimiento preventivo de equipos",
            "consultoria en gestion documental"][i % 5]
    compania, nit_c, tel, mail = EMISORES["contrato"][i % 4]
    persona = EMISORES["personas"][i % 10]
    meses = [12, 6, 3, 12, 4][i % 5]
    ini = datetime.date(2024, (i % 4) + 1, 15)
    vale = 2500000 + i * 750000
    garantia = "dos mensualidades" if i % 5 == 0 else "un deposito de $%s" % "{:,}".format(vale)
    return ("CONTRATO DE %s\n"
            "Entre %s, NIT %s, representada legalmente por %s, y la contratista %s, se celebra el presente contrato.\n\n"
            "CLÁUSULA PRIMERA - OBJETO: %s.\n\n"
            "CLÁUSULA SEGUNDA - CONTRAPRESTACION: valor mensual de COP $%s pagadero dentro de los primeros dias del mes.\n\n"
            "CLÁUSULA TERCERA - PLAZO: %d meses contados a partir del %s con renovacion automatica anual.\n\n"
            "CLÁUSULA CUARTA - GARANTIA: %s en calidad de caucion.\n\n"
            "CLÁUSULA QUINTA - TERMINACION: cualquiera de las partes podra terminar el contrato con preaviso de sesenta (60) dias.\n\n"
            "Contacto contractual: %s - Telefono %s\n\n"
            "Firmado en Bucaramanga, Santander, el %s.\n"
            ) % ("ARRENDAMIENTO" if i % 5 == 0 else tipo.upper(),
                 compania, nit_c, persona, EMISORES["personas"][(i + 3) % 10],
                 tipo, "{:,}".format(vale), meses, fecha_es(ini), garantia, mail, tel, fecha_es(ini))


def build_correspondencia(i):
    tema = ["nueva politica de trabajo hibrido",
            "mantenimiento programado del sistema de facturacion",
            "convocatoria a reunion trimestral de resultados",
            "actualizacion del procedimiento de compras",
            "notificacion de cambio de sede principal",
            "invitacion a jornada de seguridad y salud en el trabajo",
            "comunicado sobre entrega de dotacion",
            "recordatorio de uso responsable del correo corporativo"][i % 8]
    empresa = EMISORES["orgs_comun"][i % 3]
    persona = EMISORES["personas"][(i + 1) % 10]
    lugar, fecha = ("Bucaramanga", datetime.date(2024, (i % 6) + 1, 8 + i))
    return ("%s, %s\n\n"
            "Memo. No. RH-%03d\n"
            "Empresa: %s\n"
            "De: %s\n"
            "Para: Todos los empleados\n"
            "Asunto: %s\n\n"
            "Cordial saludo. Por medio del presente comunicamos que la medida descrita en el asunto\n"
            "entrara en vigencia el dia %s. Agradecemos su comprension y colaboracion.\n\n"
            "Ante cualquier inquietud, la Mesa de Ayuda atendera en el telefono +57 607 654 3210\n"
            "extension 115 y en el correo soporte@%s.\n\n"
            "Atentamente,\n%s\n"
            ) % (lugar, fecha_es(fecha), i + 1, empresa, persona, tema,
                 fecha_es(fecha + datetime.timedelta(days=7)), empresa.lower().replace(" ", ""), persona)


def build_informe(i):
    ambito = ["desempeño comercial del trimestre",
              "auditoria interna de procesos contables",
              "gestion del inventario y almacen",
              "clima organizacional y rotacion de personal",
              "cumplimiento del plan de mantenimiento",
              "estado de cartera y cobranza",
              "indicadores de calidad del servicio"][i % 7]
    empresa = EMISORES["orgs_comun"][i % 3]
    pct = 78 + (i * 3) % 20
    ventas = (84000000 + i * 5200000)
    persona = EMISORES["personas"][(i + 5) % 10]
    fecha = datetime.date(2024, (i % 4) + 2, 10 + i)
    return ("INFORME: %s\n"
            "Empresa: %s\n\n"
            "RESUMEN EJECUTIVO: el periodo analizado cerró con un cumplimiento del %d%% frente a las metas.\n\n"
            "METODOLOGIA: analisis cuantitativo y cualitativo de los indicadores definidos, con base en los\n"
            "datos recolectados por el equipo responsable durante el periodo evaluado.\n\n"
            "HALLAZGOS: se observaron mejoras en los procesos centrales; dos observaciones menores fueron\n"
            "documentadas y estan en plan de accion. No se evidenciaron irregularidades mayores.\n\n"
            "INDICADORES: ventas consolidadas $%s, cumplimiento de entregas 94.5%%, productividad general %d%%.\n\n"
            "RECOMENDACIONES: fortalecer la gestion de cobro, documentar procedimientos y capacitar al personal\n"
            "auxiliar en el uso del sistema documental.\n\n"
            "CONCLUSIONES: el desempeño del periodo es favorable y las recomendaciones permiten proyectar\n"
            "un cierre de año positivo.\n\n"
            "Elaborado por: %s - %s\n"
            ) % (ambito, empresa, pct, "{:,}".format(ventas), pct, persona, fecha_es(fecha))


def to_txt(name, text):
    with open(name, "w", encoding="utf-8") as f:
        f.write(text)


def to_pdf(name, text):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    doc = SimpleDocTemplate(name, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
    style = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14)
    stitle = ParagraphStyle("h", parent=style, fontName="Helvetica-Bold", fontSize=12, spaceAfter=10)
    story = []
    lines = text.splitlines()
    story.append(Paragraph(lines[0].strip() if lines else "", stitle))
    for ln in lines[1:]:
        story.append(Paragraph(ln.strip() or "&nbsp;", style))
        story.append(Spacer(1, 2))
    doc.build(story)


def to_docx(name, text):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    lines = text.splitlines()
    doc.add_heading(lines[0].strip() if lines else "", level=1)
    for ln in lines[1:]:
        if not ln.strip():
            continue
        p = doc.add_paragraph(ln.strip())
        p.paragraph_format.space_after = Pt(4)
    doc.save(name)


BUILDERS = {
    "factura": build_factura,
    "contrato": build_contrato,
    "correspondencia": build_correspondencia,
    "informe": build_informe,
}
COUNTS = {"factura": 9, "contrato": 9, "correspondencia": 8, "informe": 7}
FMT_CYCLE = ["txt", "pdf", "docx"]


def main(dest="../../11_Documentos_Prueba"):
    dest = os.path.abspath(dest)
    os.makedirs(dest, exist_ok=True)
    manifest = []
    counter = 0
    for cat, n in COUNTS.items():
        prefix = {"factura": "FE", "contrato": "CT", "correspondencia": "CO", "informe": "IN"}[cat]
        for i in range(n):
            text = BUILDERS[cat](i)
            ext = FMT_CYCLE[i % 3]
            name = "%s-%03d.%s" % (prefix, i + 1, ext)
            full = os.path.join(dest, name)
            if ext == "txt":
                to_txt(full, text)
            elif ext == "pdf":
                to_pdf(full, text)
            else:
                to_docx(full, text)
            primera = text.splitlines()[0] if text.splitlines() else ""
            manifest.append([name, cat, ext, primera.strip(), ""])
            counter += 1
    mani_file = os.path.join(dest, "_manifesto_corpus.csv")
    with open(mani_file, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["ID", "Archivo", "Categoria", "Formato", "Contenido (primera linea)", "Informacion esperada", "Uso en prueba"])
        for i, row in enumerate(manifest, 1):
            w.writerow([i] + row)
    print("Creados %d documentos en %s" % (counter, dest))
    print("Manifiesto: %s" % mani_file)
    return counter


if __name__ == "__main__":
    d = sys.argv[1] if len(sys.argv) > 1 else "../../11_Documentos_Prueba"
    main(d)